"""
Cover-letter support for the deterministic auto-apply flow:
- `automation/browser/selectors.py::find_cover_letter_file_input` — DOM
  detection, tested against real rendered HTML (same convention as
  test_selectors.py).
- `automation/ats/base.py::ATSAdapter.detect_cover_letter_field` /
  `find_cover_letter_text_field` / `fill_cover_letter_field` — tested via
  `GenericAdapter` against a hand-built page (same convention as
  test_generic_adapter.py), since these are concrete `ATSAdapter` methods
  every adapter inherits unchanged.
- `automation/applications/cover_letter.py` — pure unit tests (LLM call
  mocked; no browser needed).
- `ApplicationFlowManager`'s cover-letter orchestration (generate once per
  run, skip when no job description, cache/cleanup the temp file) — pure
  unit tests with a mock adapter, same convention as this file's
  `_fake_flow_manager` helper in test_application_flow_manager.py.
"""

from __future__ import annotations

from pathlib import Path
from unittest.mock import MagicMock, patch

import docx
import pytest

from app.core.crypto import encrypt_field
from app.models.db_models import CandidateProfile, ProfileDocument
from automation.applications.application_flow_manager import ApplicationFlowManager
from automation.applications.cover_letter import generate_cover_letter_text, write_cover_letter_docx
from automation.ats.generic.generic_adapter import GenericAdapter
from automation.browser.selectors import find_cover_letter_file_input


def _render(page, html: str):
    page.set_content(f"<!DOCTYPE html><html><body>{html}</body></html>")


# ---------------------------------------------------------------------------
# find_cover_letter_file_input (selectors.py)
# ---------------------------------------------------------------------------

def test_finds_a_cover_letter_file_input_alongside_a_resume_one(page):
    _render(page, '<input type="file" name="resume"><input type="file" name="cover_letter">')
    found = find_cover_letter_file_input(page)
    assert found is not None
    assert found.get_attribute("name") == "cover_letter"


def test_a_single_file_input_is_never_treated_as_a_cover_letter_field(page):
    # Even if it happens to be named "cover_letter" — a lone file input is
    # always the résumé's; see the function's own docstring.
    _render(page, '<input type="file" name="cover_letter">')
    assert find_cover_letter_file_input(page) is None


def test_returns_none_when_no_input_matches_a_cover_letter_hint(page):
    _render(page, '<input type="file" name="resume"><input type="file" name="additional_document">')
    assert find_cover_letter_file_input(page) is None


# ---------------------------------------------------------------------------
# ATSAdapter.detect_cover_letter_field / fill_cover_letter_field (base.py),
# exercised through GenericAdapter — concrete base-class methods every
# adapter inherits unchanged.
# ---------------------------------------------------------------------------

def _profile(**overrides) -> CandidateProfile:
    defaults = dict(profile_id="profile-1", user_id="user-1", full_name="Ada Lovelace", email="ada@example.com")
    defaults.update(overrides)
    profile = CandidateProfile(**defaults)
    profile.phone_encrypted = encrypt_field("+1-555-0100")
    return profile


def _resume_document(path) -> ProfileDocument:
    return ProfileDocument(
        document_id="doc-1", profile_id="profile-1", document_type="resume",
        original_filename="resume.pdf", stored_path=str(path), file_hash="abc123", is_default=True,
    )


@pytest.fixture
def resume_file(tmp_path):
    path = tmp_path / "resume.pdf"
    path.write_bytes(b"%PDF-1.4 fake resume bytes")
    return path


def test_detect_returns_none_none_when_the_page_asks_for_no_cover_letter(page, resume_file):
    _render(page, '<input type="text" name="email">')
    adapter = GenericAdapter(page=page, profile=_profile(), resume_document=_resume_document(resume_file))
    assert adapter.detect_cover_letter_field() == (None, None)


def test_detect_finds_a_free_text_cover_letter_box(page, resume_file):
    _render(page, '<label for="cl">Cover Letter</label><textarea id="cl"></textarea>')
    adapter = GenericAdapter(page=page, profile=_profile(), resume_document=_resume_document(resume_file))
    kind, locator = adapter.detect_cover_letter_field()
    assert kind == "text"
    assert locator.get_attribute("id") == "cl"


def test_detect_prefers_text_over_a_coexisting_file_input(page, resume_file):
    _render(
        page,
        """
        <label for="cl">Cover Letter</label><textarea id="cl"></textarea>
        <input type="file" name="resume"><input type="file" name="cover_letter">
        """,
    )
    adapter = GenericAdapter(page=page, profile=_profile(), resume_document=_resume_document(resume_file))
    kind, _locator = adapter.detect_cover_letter_field()
    assert kind == "text"


def test_detect_finds_a_dedicated_cover_letter_file_input(page, resume_file):
    _render(page, '<input type="file" name="resume"><input type="file" name="cover_letter">')
    adapter = GenericAdapter(page=page, profile=_profile(), resume_document=_resume_document(resume_file))
    kind, locator = adapter.detect_cover_letter_field()
    assert kind == "file"
    assert locator.get_attribute("name") == "cover_letter"


def test_fill_cover_letter_field_types_prose_into_a_text_box(page, resume_file):
    _render(page, '<label for="cl">Cover Letter</label><textarea id="cl"></textarea>')
    adapter = GenericAdapter(page=page, profile=_profile(), resume_document=_resume_document(resume_file))
    kind, locator = adapter.detect_cover_letter_field()
    result = adapter.fill_cover_letter_field(locator, "Dear Hiring Manager, I am excited to apply.")
    assert result.filled
    assert result.field_key == "cover_letter"
    assert page.locator("#cl").input_value() == "Dear Hiring Manager, I am excited to apply."


def test_fill_cover_letter_field_marks_the_element_examined_so_it_is_not_refilled(page, resume_file):
    _render(page, '<label for="cl">Cover Letter</label><textarea id="cl"></textarea>')
    adapter = GenericAdapter(page=page, profile=_profile(), resume_document=_resume_document(resume_file))
    kind, locator = adapter.detect_cover_letter_field()
    adapter.fill_cover_letter_field(locator, "Some generated letter text.")
    assert page.locator("#cl").get_attribute("data-automation-examined") == "1"


def test_fill_cover_letter_field_uploads_a_file(page, resume_file, tmp_path):
    letter_path = tmp_path / "cover_letter.docx"
    letter_path.write_bytes(b"fake docx bytes")
    _render(page, '<input type="file" name="resume"><input type="file" name="cover_letter">')
    adapter = GenericAdapter(page=page, profile=_profile(), resume_document=_resume_document(resume_file))
    kind, locator = adapter.detect_cover_letter_field()
    assert kind == "file"
    result = adapter.fill_cover_letter_field(locator, str(letter_path))
    assert result.filled
    assert result.value_used == "cover_letter.docx"


# ---------------------------------------------------------------------------
# automation/applications/cover_letter.py — pure unit tests, no browser
# ---------------------------------------------------------------------------

def test_generate_cover_letter_text_calls_the_llm_router_with_job_description_and_profile_facts():
    profile = _profile(current_role="Backend Engineer", current_company="Acme", years_of_experience=5.0)
    with patch("automation.applications.cover_letter.generate_answer", return_value="  Dear Hiring Manager...  ") as mock_generate:
        text = generate_cover_letter_text(profile=profile, job_description="We need a backend engineer.", company="Acme", position="Backend Engineer")
    assert text == "Dear Hiring Manager..."  # stripped
    mock_generate.assert_called_once()
    kwargs = mock_generate.call_args.kwargs
    assert kwargs["task"] == "cover_letter_generation"
    assert "We need a backend engineer." in kwargs["prompt"]
    assert "Backend Engineer" in kwargs["prompt"]
    assert "Acme" in kwargs["prompt"]


def test_write_cover_letter_docx_produces_a_readable_file_with_the_given_text():
    path = write_cover_letter_docx("First paragraph.\n\nSecond paragraph.")
    try:
        document = docx.Document(path)
        full_text = "\n".join(p.text for p in document.paragraphs)
        assert "First paragraph." in full_text
        assert "Second paragraph." in full_text
    finally:
        Path(path).unlink(missing_ok=True)


# ---------------------------------------------------------------------------
# ApplicationFlowManager orchestration — mock adapter, no real browser.
# ---------------------------------------------------------------------------

def _manager(**overrides) -> ApplicationFlowManager:
    defaults = dict(
        application_id="app-1",
        user_id="user-1",
        job_url="data:text/html,<html></html>",
        ats_platform="greenhouse",
        adapter_cls=object,  # never instantiated in these tests
        profile=_profile(),
        resume_document=None,
        browser_manager=MagicMock(),
    )
    defaults.update(overrides)
    return ApplicationFlowManager(**defaults)


def test_does_nothing_when_the_adapter_detects_no_cover_letter_field():
    manager = _manager(job_description="A real job description.")
    adapter = MagicMock()
    adapter.detect_cover_letter_field.return_value = (None, None)
    all_results = []

    manager._handle_cover_letter_if_requested(adapter, page_index=0, all_results=all_results)

    adapter.fill_cover_letter_field.assert_not_called()
    assert all_results == []


def test_generates_and_fills_a_text_field_when_one_is_requested():
    manager = _manager(job_description="A real job description.")
    adapter = MagicMock()
    adapter.detect_cover_letter_field.return_value = ("text", "the-locator")
    fake_result = MagicMock(filled=True)
    adapter.fill_cover_letter_field.return_value = fake_result
    all_results = []

    with patch("automation.applications.application_flow_manager.generate_cover_letter_text", return_value="Generated letter.") as mock_generate:
        manager._handle_cover_letter_if_requested(adapter, page_index=0, all_results=all_results)

    mock_generate.assert_called_once_with(profile=manager.profile, job_description="A real job description.")
    adapter.fill_cover_letter_field.assert_called_once_with("the-locator", "Generated letter.")
    assert all_results == [fake_result]


def test_generation_happens_at_most_once_across_multiple_pages():
    manager = _manager(job_description="A real job description.")
    adapter = MagicMock()
    adapter.detect_cover_letter_field.return_value = ("text", "the-locator")
    adapter.fill_cover_letter_field.return_value = MagicMock(filled=True)

    with patch("automation.applications.application_flow_manager.generate_cover_letter_text", return_value="Generated letter.") as mock_generate:
        manager._handle_cover_letter_if_requested(adapter, page_index=0, all_results=[])
        manager._handle_cover_letter_if_requested(adapter, page_index=1, all_results=[])

    mock_generate.assert_called_once()
    assert adapter.fill_cover_letter_field.call_count == 2  # filled again on page 2, but with the SAME cached text
    assert adapter.fill_cover_letter_field.call_args_list[1].args[1] == "Generated letter."


def test_no_job_description_means_the_field_is_left_unfilled_not_invented():
    manager = _manager(job_description=None)
    adapter = MagicMock()
    adapter.detect_cover_letter_field.return_value = ("text", "the-locator")
    all_results = []

    manager._handle_cover_letter_if_requested(adapter, page_index=0, all_results=all_results)

    adapter.fill_cover_letter_field.assert_not_called()
    assert all_results == []


def test_a_failed_generation_leaves_the_field_unfilled_rather_than_failing_the_run():
    manager = _manager(job_description="A real job description.")
    adapter = MagicMock()
    adapter.detect_cover_letter_field.return_value = ("text", "the-locator")

    with patch("automation.applications.application_flow_manager.generate_cover_letter_text", side_effect=RuntimeError("LLM down")):
        manager._handle_cover_letter_if_requested(adapter, page_index=0, all_results=[])  # must not raise

    adapter.fill_cover_letter_field.assert_not_called()


def test_a_file_request_materializes_a_temp_file_and_reuses_it_on_a_later_page(tmp_path):
    manager = _manager(job_description="A real job description.")
    adapter = MagicMock()
    adapter.detect_cover_letter_field.return_value = ("file", "the-locator")
    adapter.fill_cover_letter_field.return_value = MagicMock(filled=True)
    fake_path = str(tmp_path / "generated.docx")
    Path(fake_path).write_bytes(b"fake docx")

    with patch("automation.applications.application_flow_manager.generate_cover_letter_text", return_value="Generated letter."), \
         patch("automation.applications.application_flow_manager.write_cover_letter_docx", return_value=fake_path) as mock_write:
        manager._handle_cover_letter_if_requested(adapter, page_index=0, all_results=[])
        manager._handle_cover_letter_if_requested(adapter, page_index=1, all_results=[])

    mock_write.assert_called_once()  # not regenerated for the second page
    assert adapter.fill_cover_letter_field.call_args_list[0].args[1] == fake_path
    assert adapter.fill_cover_letter_field.call_args_list[1].args[1] == fake_path


def test_cleanup_deletes_the_temp_file_and_never_touches_the_resume_document(tmp_path):
    manager = _manager()
    temp_path = tmp_path / "generated.docx"
    temp_path.write_bytes(b"fake docx")
    manager._cover_letter_temp_path = str(temp_path)

    manager._cleanup_cover_letter_file()

    assert not temp_path.exists()
    assert manager._cover_letter_temp_path is None
    assert manager.resume_document is None  # untouched — this test never gave it one to begin with


def test_cleanup_is_a_no_op_when_no_file_was_ever_created():
    manager = _manager()
    manager._cleanup_cover_letter_file()  # must not raise
    assert manager._cover_letter_temp_path is None


def test_finish_browser_session_cleans_up_the_temp_cover_letter_file_when_closing():
    manager = _manager()
    manager.browser_manager.headless = True
    with patch.object(manager, "_cleanup_cover_letter_file") as mock_cleanup:
        manager._finish_browser_session(context=MagicMock(), status="applied")
    mock_cleanup.assert_called_once()


def test_finish_browser_session_cleans_up_even_when_the_browser_stays_open_for_review():
    manager = _manager()
    manager.browser_manager.headless = False
    try:
        with patch.object(manager, "_cleanup_cover_letter_file") as mock_cleanup:
            manager._finish_browser_session(context=MagicMock(), status="copilot_review")
        mock_cleanup.assert_called_once()
    finally:
        from automation.applications.application_flow_manager import _OPEN_REVIEW_SESSIONS
        _OPEN_REVIEW_SESSIONS.pop(manager.application_id, None)
